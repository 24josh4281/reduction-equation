#!/usr/bin/env python3
"""
PPT Master - Reduction-Project Monitoring Plan (rough text -> editable equations + tables)

One plan file holds a reduction project's formulas (typed as rough text, see
`math_shorthand.py`) and its parameter dictionary. From it this tool derives
everything the KCS methodology / business-plan forms repeat by hand:

    - numbered display equations                         (native OMML)
    - a "기호 | 정의 | 단위 (| 값)" table under each section  (symbols as OMML)
    - the monitoring summary table  "구분 | 기호 | 항목 | 고정/모니터링 | 비고"
    - one "데이터/인자" detail table per fixed and per monitored parameter

Symbols are matched automatically: every symbol used in an equation is looked
up in the parameter dictionary; missing ones are added with a guessed kind
(left-hand side -> 산출, a subscript holding y -> 모니터링, else 고정) and
reported, so nothing silently drops out of the tables.

Quick text (the form typed in the web editor or a .txt file):
    # 베이스라인 배출량               new section (## for a sub-section)
    > 설명 문장                       section description
    BEy = BE_MAT,y + BE_ENE,y        equation (optional "(식 1)" at the end)
    Q_PJ,j,y | y년도 사업필터 순판매량 | 개/년 | 모니터링 | 값
                                     parameter row: symbol | 정의 | 단위 | 구분 | 값
Inside any text field, $...$ is inline math (shorthand or LaTeX).

Subcommands:
    example                         print a sample plan (JSON)
    from-text <in.txt> -o plan.json  quick text -> plan
    import <file.docx|pptx> -o plan.json
                                    read equations and parameter tables of an
                                    existing methodology / plan / deck
    check <plan.json>               derive and print issues (exit 1 on errors)
    docx <plan.json> -o out.docx    KCS-style Word section (equations + tables)
    pptx <plan.json> -o out.pptx    16:9 slides (equations + tables)
    spec <plan.json> -o spec.json   office_equations.py spec (for inject-docx /
                                    inject-pptx tokens {{eq:id}} {{sym:id}})

Usage:
    python3 scripts/monitoring_plan.py import 방법론.docx -o projects/x/plan.json
    python3 scripts/monitoring_plan.py docx projects/x/plan.json -o projects/x/exports/모니터링계획.docx
    python3 scripts/monitoring_plan.py pptx projects/x/plan.json -o projects/x/exports/수식_인자.pptx

Dependencies:
    python-pptx (pptx subcommand only)
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree as ET

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from math_shorthand import split_label, to_latex, to_shorthand  # noqa: E402
from omml_math import (  # noqa: E402
    GREEK,
    Delim,
    EquationError,
    Frac,
    Nary,
    Rad,
    Run,
    Script,
    Accent,
    _is_hangul,
    linear_text,
    mathml,
    parse_latex,
)

KINDS = {"calculated": "산출", "monitoring": "모니터링", "fixed": "고정"}
KIND_ALIASES = {"산출": "calculated", "산출값": "calculated", "계산": "calculated",
                "모니터링": "monitoring", "측정": "monitoring", "monitoring": "monitoring",
                "고정": "fixed", "고정값": "fixed", "사전결정": "fixed", "사전 결정": "fixed", "fixed": "fixed", "calculated": "calculated"}
# KCS "데이터/인자" table rows, in form order.
DETAIL_FIELDS = [("unit", "데이터 단위"), ("name", "설명"), ("source", "데이터 출처"),
                 ("value", "적용된 값"), ("method", "측정 방법/절차"), ("frequency", "모니터링 주기"),
                 ("qaqc", "QA/QC 절차"), ("purpose", "데이터 목적"), ("note", "기타 의견")]
# Submission forms. KCS follows the 뉴라이즌 KCS submissions (v1.1/v1.2, 2026-09);
# the 환경부 외부사업 labels follow the CDM-derived 외부사업 methodology layout and
# must be checked against the current official form before submission.
# plan["options"]["labels"] overrides any entry for one plan.
FORMS: dict[str, dict[str, Any]] = {
    "kcs": {
        "label": "대한상의 탄소감축인증센터(KCS)",
        "symbol_style": "table",
        "param_word": "데이터/인자",
        "chapter": "모니터링 방법론",
        "procedure": "모니터링 절차",
        "procedure_lead": "사업기간 동안 베이스라인 배출량, 사업 배출량, 누출량을 산정하기 위하여 "
                          "모니터링해야 하는 인자들은 다음과 같다.",
        "summary": [("no", "구분", 700), ("symbol", "기호", 1900), ("name", "항목", 4300),
                    ("kind", "고정/모니터링", 1400), ("remark", "비고", 1200)],
        "kind_names": {"fixed": "고정", "monitoring": "모니터링"},
        "groups": [("fixed", "고정 데이터 및 인자",
                    "사업기간 동안 모니터링하지 않아도 되는 고정 값을 가지는 인자들은 아래와 같다."),
                   ("monitoring", "모니터링 데이터 및 인자",
                    "사업기간 동안 모니터링해야 하는 데이터 및 인자들은 아래와 같다.")],
        "fields": {"fixed": DETAIL_FIELDS, "monitoring": DETAIL_FIELDS},
    },
    "moe": {
        "label": "환경부 외부사업 방법론",
        "symbol_style": "where",
        "param_word": "데이터/변수",
        "chapter": "모니터링 계획",
        "procedure": "모니터링 대상 데이터 및 변수 목록",
        "procedure_lead": "베이스라인 배출량, 사업 배출량 및 누출량 산정에 필요한 데이터와 변수는 다음과 같다.",
        "summary": [("no", "번호", 700), ("symbol", "데이터/변수", 1900), ("name", "설명", 3900),
                    ("unit", "단위", 1500), ("kind", "구분", 1500)],
        "kind_names": {"fixed": "사전 결정", "monitoring": "모니터링"},
        "groups": [("fixed", "사전에 결정되는 데이터 및 변수",
                    "사업기간 동안 모니터링하지 않고 사전에 결정되는 데이터와 변수는 다음과 같다."),
                   ("monitoring", "모니터링 대상 데이터 및 변수",
                    "사업기간 동안 모니터링해야 하는 데이터와 변수는 다음과 같다.")],
        "fields": {
            "fixed": [("unit", "데이터 단위"), ("name", "설명"), ("source", "데이터 출처"),
                      ("value", "적용된 값"), ("method", "측정 방법 및 절차"), ("purpose", "데이터 목적"),
                      ("note", "기타 의견")],
            "monitoring": [("unit", "데이터 단위"), ("name", "설명"), ("source", "데이터 출처"),
                           ("value", "적용된 값"), ("method", "측정 방법 및 절차"),
                           ("frequency", "모니터링 주기"), ("qaqc", "QA/QC 절차"), ("purpose", "데이터 목적"),
                           ("note", "기타 의견")],
        },
    },
}


def form_of(options: dict[str, Any]) -> dict[str, Any]:
    """Form profile for options["form"] ("kcs" | "moe") with options["labels"] overrides."""
    base = FORMS.get(options.get("form") or "kcs", FORMS["kcs"])
    return {**base, **(options.get("labels") or {})}


PARAM_FIELDS = ["symbol", "name", "unit", "kind", "value", "source", "method", "frequency",
                "qaqc", "purpose", "note", "remark"]
INLINE_MATH = re.compile(r"\$([^$\n]+)\$")
GREEK_NAME = {v: k for k, v in GREEK.items() if not k.startswith("var")}
# Emission totals are computed, never measured: BE_y, PE_{MAT,y}, LE_y, ER_y.
CALCULATED_BASES = {"BE", "PE", "LE", "ER"}


# ---------------------------------------------------------------------------
# Symbols
# ---------------------------------------------------------------------------

def _flat(nodes: list) -> str:
    """Compact text of a subscript ("PJ,j,y")."""
    from omml_math import _linear
    out = []
    for n in nodes:
        if isinstance(n, Run):
            out.append(n.text)
        elif isinstance(n, Script):
            out.append(_flat(n.base) + ("_" + _flat(n.sub) if n.sub else ""))
        else:
            out.append(_linear(n))
    return "".join(out).replace(" ", "")


def symbol_key(base: str, sub: Optional[str] = None, bar: bool = False) -> str:
    """Canonical LaTeX of one symbol: BE_{MFG,y}, \\sigma, \\overline{DP}_{BL,a}."""
    base = f"\\{GREEK_NAME[base]}" if base in GREEK_NAME else base
    if bar:
        base = f"\\overline{{{base}}}"
    return f"{base}_{{{sub}}}" if sub else base


def _symbol_base(nodes: list) -> Optional[tuple[str, bool]]:
    """(letters, overlined) when nodes are one variable, optionally under a bar."""
    if len(nodes) == 1 and isinstance(nodes[0], Run) and nodes[0].kind == "var":
        return nodes[0].text, False
    if len(nodes) == 1 and isinstance(nodes[0], Accent) and nodes[0].bar:
        inner = _symbol_base(nodes[0].body)
        return (inner[0], True) if inner and not inner[1] else None
    return None


def _walk_symbols(nodes: list, out: list[str], indices: set[str]) -> None:
    for n in nodes:
        if isinstance(n, Run):
            if n.kind == "var":
                out.append(symbol_key(n.text))
        elif isinstance(n, Script):
            sym = _symbol_base(n.base)
            if sym:
                out.append(symbol_key(sym[0], _flat(n.sub) if n.sub else None, sym[1]))
            else:
                _walk_symbols(n.base, out, indices)
            if n.sup:
                _walk_symbols(n.sup, out, indices)
        elif isinstance(n, Nary):
            for part in (n.sub or []):
                if isinstance(part, Run) and part.kind in ("var", "label"):
                    indices.add(part.text)
            if n.sup:
                _walk_symbols(n.sup, out, indices)
            _walk_symbols(n.body, out, indices)
        elif isinstance(n, Frac):
            _walk_symbols(n.num, out, indices)
            _walk_symbols(n.den, out, indices)
        elif isinstance(n, Accent) and _symbol_base([n]):
            letters, bar = _symbol_base([n]) or ("", False)
            out.append(symbol_key(letters, None, bar))
        elif isinstance(n, (Rad, Delim, Accent)):
            _walk_symbols(n.body, out, indices)


def equation_symbols(latex: str) -> tuple[list[str], list[str], set[str]]:
    """Return (lhs symbols, all symbols in order, n-ary index letters)."""
    nodes = parse_latex(latex)
    split = next((i for i, n in enumerate(nodes) if isinstance(n, Run) and n.kind == "op"
                  and n.text == "="), None)
    lhs: list[str] = []
    every: list[str] = []
    indices: set[str] = set()
    if split is not None:
        _walk_symbols(nodes[:split], lhs, indices)
    _walk_symbols(nodes, every, indices)
    return lhs, list(dict.fromkeys(every)), indices


def param_key(symbol: str) -> str:
    """Canonical key (LaTeX) of a parameter symbol typed as shorthand or LaTeX."""
    _, syms, _ = equation_symbols(to_latex(symbol))
    if len(syms) != 1:
        raise EquationError(f"parameter symbol must be one symbol: {symbol!r}")
    return syms[0]


def key_slug(key: str) -> str:
    """Token-safe id ({{sym:Q_PJ.j.y}}, {{sym:DPbar_BL.a}})."""
    slug = re.sub(r"\\overline\{([^}]*)\}", r"\1bar", key)
    slug = slug.replace("\\", "").replace("{", "").replace("}", "").replace(",", ".")
    return re.sub(r"[^A-Za-z0-9_.-]", "", slug) or "sym"


# ---------------------------------------------------------------------------
# Quick text <-> plan
# ---------------------------------------------------------------------------

def new_plan() -> dict[str, Any]:
    return {"title": "", "font": "Pretendard", "font_size_pt": 10.5,
            "options": {"form": "kcs", "number_equations": True, "label_prefix": "식 ", "show_values": False},
            "sections": [], "parameters": []}


def parse_param_line(line: str) -> dict[str, Any]:
    cells = [c.strip() for c in line.split("|")]
    order = ["symbol", "name", "unit", "kind", "value", "source", "method", "frequency",
             "qaqc", "purpose", "note"]
    param = {key: cells[i] for i, key in enumerate(order) if i < len(cells) and cells[i]}
    if "kind" in param:
        param["kind"] = KIND_ALIASES.get(param["kind"], param["kind"])
    return param


def plan_from_text(text: str, plan: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """Rebuild sections/equations from quick text; parameter rows are merged into the dictionary."""
    plan = plan or new_plan()
    sections: list[dict[str, Any]] = []
    current: Optional[dict[str, Any]] = None
    eq_no = 0
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("//"):
            continue
        if line.startswith("#"):
            level = 2 if line.startswith("##") else 1
            current = {"id": f"s{len(sections) + 1}", "name": line.lstrip("#").strip(),
                       "level": level, "description": "", "equations": []}
            sections.append(current)
            continue
        if line.startswith(">"):
            if current is not None:
                current["description"] = (current["description"] + " " + line[1:].strip()).strip()
            continue
        if "|" in line:
            merge_parameter(plan, parse_param_line(line))
            continue
        if current is None:
            current = {"id": "s1", "name": "산정식", "level": 1, "description": "", "equations": []}
            sections.append(current)
        eq_no += 1
        formula, label = split_label(line)
        current["equations"].append({"id": f"e{eq_no}", "text": formula, "label": label or ""})
    plan["sections"] = sections
    return plan


def plan_to_text(plan: dict[str, Any]) -> str:
    lines: list[str] = []
    for section in plan.get("sections", []):
        if lines:
            lines.append("")
        lines.append(("## " if section.get("level") == 2 else "# ") + section.get("name", ""))
        if section.get("description"):
            lines.append("> " + section["description"])
        for eq in section.get("equations", []):
            label = eq.get("label")
            lines.append(eq.get("text", "") + (f" ({label})" if label and not plan.get("options", {})
                                                .get("number_equations", True) else ""))
    return "\n".join(lines) + "\n"


def merge_parameter(plan: dict[str, Any], incoming: dict[str, Any], overwrite: bool = True) -> None:
    if not incoming.get("symbol"):
        return
    try:
        key = param_key(incoming["symbol"])
    except EquationError:
        key = incoming["symbol"]
    for param in plan.setdefault("parameters", []):
        try:
            same = param_key(param["symbol"]) == key
        except EquationError:
            same = param["symbol"] == incoming["symbol"]
        if same:
            for field, value in incoming.items():
                if field == "symbol" or value in (None, ""):
                    continue
                if overwrite or not param.get(field):
                    param[field] = value
            param.pop("auto", None)
            return
    plan["parameters"].append({k: v for k, v in incoming.items() if v not in (None, "")})


# ---------------------------------------------------------------------------
# Derivation
# ---------------------------------------------------------------------------

def _purpose_for(section_names: list[str]) -> str:
    found: list[str] = []
    for name in section_names:
        for word, label in (("베이스라인", "베이스라인 배출량 산정"), ("누출", "누출량 산정"),
                            ("사업", "사업 배출량 산정"), ("감축", "감축량 산정")):
            if word in name:
                if label not in found:
                    found.append(label)
                break
    return ", ".join(found)


def _guess_kind(key: str, lhs: set[str]) -> str:
    base, _, sub = key.partition("_")
    if key in lhs or base in CALCULATED_BASES:
        return "calculated"
    return "monitoring" if "y" in sub.strip("{}").split(",") else "fixed"


def derive(plan: dict[str, Any]) -> dict[str, Any]:
    """Parse every equation, match symbols to parameters and list issues."""
    issues: list[dict[str, str]] = []
    opts = {**new_plan()["options"], **plan.get("options", {})}
    params = plan.get("parameters", [])
    by_key: dict[str, dict[str, Any]] = {}
    for idx, param in enumerate(params):
        try:
            key = param_key(param.get("symbol", ""))
        except EquationError as exc:
            issues.append({"level": "error", "where": f"인자 {idx + 1}", "msg": f"기호 해석 실패: {exc}"})
            continue
        if key in by_key:
            issues.append({"level": "warn", "where": f"인자 {param.get('symbol')}", "msg": "같은 기호가 두 번 정의됨"})
        by_key[key] = param

    lhs_all: set[str] = set()
    defined_in: dict[str, str] = {}
    used_in: dict[str, list[str]] = {}
    sections_out: list[dict[str, Any]] = []
    counter = 0
    for s_idx, section in enumerate(plan.get("sections", []), start=1):
        sec_out = {**section, "equations": [], "symbols": []}
        sec_syms: list[str] = []
        for eq in section.get("equations", []):
            out = dict(eq)
            text = eq.get("text") or eq.get("latex") or ""
            try:
                latex = eq.get("latex") if not eq.get("text") else to_latex(text)
                lhs, syms, indices = equation_symbols(latex)
                out.update(latex=latex, mathml=mathml(latex), linear=linear_text(latex),
                           lhs=lhs, symbols=syms)
                for key in lhs:
                    if key in defined_in:
                        issues.append({"level": "warn", "where": f"{section.get('name')}",
                                       "msg": f"{key}가 {defined_in[key]}에서도 정의됨"})
                    defined_in.setdefault(key, section.get("name", ""))
                    lhs_all.add(key)
                for key in syms:
                    if key in indices:
                        continue
                    if key not in sec_syms:
                        sec_syms.append(key)
                    used_in.setdefault(key, [])
                    if section.get("name", "") not in used_in[key]:
                        used_in[key].append(section.get("name", ""))
            except EquationError as exc:
                out.update(latex=None, error=str(exc))
                issues.append({"level": "error", "where": f"{section.get('name')}: {text[:40]}",
                               "msg": f"수식 해석 실패: {exc}"})
            if opts.get("number_equations") and not out.get("error"):
                counter += 1
                if not out.get("label"):
                    out["label"] = f"{opts.get('label_prefix', '식 ')}{counter}"
            sec_out["equations"].append(out)
        sec_out["symbols"] = sec_syms
        sections_out.append(sec_out)

    # Section purposes: a symbol's purpose follows the top-level section it serves.
    top_name = ""
    owner: dict[str, list[str]] = {}
    for sec in sections_out:
        if sec.get("level", 1) == 1:
            top_name = sec.get("name", "")
        for key in sec["symbols"]:
            owner.setdefault(key, [])
            for name in (top_name, sec.get("name", "")):
                if name and name not in owner[key]:
                    owner[key].append(name)

    rows: list[dict[str, Any]] = []
    for key in list(dict.fromkeys([k for s in sections_out for k in s["symbols"]] + list(by_key))):
        param = by_key.get(key)
        row: dict[str, Any] = {f: "" for f in PARAM_FIELDS}
        if param:
            row.update({k: v for k, v in param.items() if v is not None})
        else:
            row["symbol"] = to_shorthand(key)
            row["auto"] = True
        row["key"] = key
        row["latex"] = key
        row["id"] = key_slug(key)
        row["used"] = key in used_in
        row["kind_auto"] = not row.get("kind")
        if not row.get("kind"):
            row["kind"] = _guess_kind(key, lhs_all)
        else:
            row["kind"] = KIND_ALIASES.get(row["kind"], row["kind"])
        if not row.get("purpose") and row["kind"] != "calculated":
            row["purpose_auto"] = _purpose_for(owner.get(key, []))
        try:
            row["mathml"] = mathml(row["latex"], display=False)
        except EquationError:
            row["mathml"] = ""
        rows.append(row)
        where = f"인자 {row['symbol']}"
        if not row["used"]:
            issues.append({"level": "info", "where": where, "msg": "수식에서 쓰이지 않는 인자"})
        if not row.get("name"):
            issues.append({"level": "warn", "where": where, "msg": "정의(설명) 없음"})
        if not row.get("unit"):
            issues.append({"level": "warn", "where": where, "msg": "단위 없음"})
        if row["kind"] in ("monitoring", "fixed"):
            missing = [label for field, label in (("source", "데이터 출처"), ("method", "측정 방법"),
                                                  ("frequency", "모니터링 주기"))
                       if not row.get(field)]
            if missing:
                issues.append({"level": "warn", "where": where, "msg": "세부표 빈칸: " + ", ".join(missing)})
        if row["kind"] == "calculated" and key not in lhs_all:
            issues.append({"level": "warn", "where": where, "msg": "산출 인자인데 정의하는 수식이 없음"})
    return {"sections": sections_out, "parameters": rows, "issues": issues, "options": opts,
            "title": plan.get("title", ""), "font": plan.get("font") or "Pretendard",
            "font_size_pt": float(plan.get("font_size_pt") or 10.5)}


def to_equation_spec(derived: dict[str, Any]) -> dict[str, Any]:
    """office_equations.py spec; ids double as {{eq:id}} / {{sym:id}} tokens."""
    by_key = {p["key"]: p for p in derived["parameters"]}
    scenarios = []
    for sec in derived["sections"]:
        scenarios.append({
            "id": sec["id"], "name": sec.get("name", ""), "description": sec.get("description", ""),
            "equations": [{"id": e["id"], "label": e.get("label", ""), "latex": e["latex"]}
                          for e in sec["equations"] if e.get("latex")],
            "where": [{"symbol": by_key[k]["latex"], "desc": by_key[k].get("name", ""),
                       "unit": by_key[k].get("unit", "")} for k in sec["symbols"] if k in by_key],
        })

    def pack(p: dict[str, Any]) -> dict[str, Any]:
        return {"id": p["id"], "symbol": p["latex"], "name": p.get("name", ""), "unit": p.get("unit", ""),
                "source": p.get("source", ""), "method": p.get("method", ""),
                "frequency": p.get("frequency", ""), "qaqc": p.get("qaqc", ""), "value": p.get("value", "")}
    params = derived["parameters"]
    return {"title": derived.get("title", ""), "font": derived["font"], "font_size_pt": derived["font_size_pt"],
            "scenarios": scenarios,
            "monitoring_parameters": [pack(p) for p in params if p["kind"] == "monitoring"],
            "fixed_parameters": [pack(p) for p in params if p["kind"] == "fixed"],
            "calculated_parameters": [pack(p) for p in params if p["kind"] == "calculated"]}


# ---------------------------------------------------------------------------
# Import from an existing Word / PowerPoint file
# ---------------------------------------------------------------------------

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OP_LATEX = {"×": " \\times ", "∙": " \\cdot ", "·": " \\cdot ", "⋅": " \\cdot ", "−": " - ",
            "≤": " \\le ", "≥": " \\ge ", "≠": " \\neq ", "≈": " \\approx ", "±": " \\pm ",
            "=": " = ", "+": " + ", "-": " - ", "→": " \\rightarrow ", "∑": "\\sum ",
            "%": "\\%", "{": "\\{", "}": "\\}", "#": "\\#", "&": "\\&", "_": "\\_"}
NARY_LATEX = {"∑": "\\sum", "∏": "\\prod", "∫": "\\int"}


def _mtext(text: str) -> str:
    from math_shorthand import _normalize
    text = _normalize(text)
    out: list[str] = []
    hangul: list[str] = []

    def flush() -> None:
        if hangul:
            out.append("\\text{" + "".join(hangul).strip() + "}")
            hangul.clear()
    for ch in text:
        if _is_hangul(ch) or (hangul and ch == " "):
            hangul.append(ch)
            continue
        flush()
        if ch in GREEK_NAME:
            out.append(f"\\{GREEK_NAME[ch]} ")
        elif ch in OP_LATEX:
            out.append(OP_LATEX[ch])
        else:
            out.append(ch)
    flush()
    return "".join(out)


def _ochild(el: ET.Element, name: str) -> Optional[ET.Element]:
    return el.find(M + name)


def _oseq(el: Optional[ET.Element]) -> str:
    if el is None:
        return ""
    return "".join(omml_to_latex(c) for c in el)


def _prop(el: ET.Element, path: str, default: str) -> str:
    found = el.find(path)
    if found is None:
        return default
    return found.get(M + "val", default)


DELIM_OPEN = {"(": "\\left(", "[": "\\left[", "{": "\\left\\{", "|": "\\left|", "": "\\left."}
DELIM_CLOSE = {")": "\\right)", "]": "\\right]", "}": "\\right\\}", "|": "\\right|", "": "\\right."}


def omml_to_latex(el: ET.Element) -> str:
    tag = el.tag.replace(M, "")
    if tag in ("oMath", "e", "num", "den", "sub", "sup", "deg", "fName", "lim"):
        return _oseq(el)
    if tag == "oMathPara":
        return " ".join(omml_to_latex(c) for c in el.findall(M + "oMath"))
    if tag == "r":
        text = "".join(t.text or "" for t in el.findall(M + "t"))
        return _mtext(text)
    if tag == "sSub":
        base, sub = _oseq(_ochild(el, "e")).strip(), _oseq(_ochild(el, "sub")).strip()
        # Word sometimes puts "PE_y − LE" inside the base of LE_y; keep only the last symbol there.
        split = re.match(r"^(.*[\s=+\-(])([A-Za-z]+)$", base, re.S)
        if split:
            return split.group(1) + split.group(2) + "_{" + sub + "}"
        return "{" + base + "}_{" + sub + "}"
    if tag == "sSup":
        return "{" + _oseq(_ochild(el, "e")).strip() + "}^{" + _oseq(_ochild(el, "sup")).strip() + "}"
    if tag == "sSubSup":
        return ("{" + _oseq(_ochild(el, "e")).strip() + "}_{" + _oseq(_ochild(el, "sub")).strip()
                + "}^{" + _oseq(_ochild(el, "sup")).strip() + "}")
    if tag == "f":
        return "\\frac{" + _oseq(_ochild(el, "num")) + "}{" + _oseq(_ochild(el, "den")) + "}"
    if tag == "nary":
        pr = _ochild(el, "naryPr")
        chr_ = _prop(pr, M + "chr", "∫") if pr is not None else "∫"
        out = NARY_LATEX.get(chr_, "\\sum")
        sub, sup = _oseq(_ochild(el, "sub")).strip(), _oseq(_ochild(el, "sup")).strip()
        if sub:
            out += "_{" + sub + "}"
        if sup:
            out += "^{" + sup + "}"
        return out + " " + _oseq(_ochild(el, "e"))
    if tag == "d":
        pr = _ochild(el, "dPr")
        beg = _prop(pr, M + "begChr", "(") if pr is not None else "("
        end = _prop(pr, M + "endChr", ")") if pr is not None else ")"
        sep = _prop(pr, M + "sepChr", "|") if pr is not None else "|"
        parts = [_oseq(e) for e in el.findall(M + "e")]
        inner = (f" {sep} " if sep != "," else ", ").join(parts)
        return DELIM_OPEN.get(beg, "\\left(") + inner + DELIM_CLOSE.get(end, "\\right)")
    if tag == "rad":
        deg = _oseq(_ochild(el, "deg")).strip()
        return ("\\sqrt[" + deg + "]{" if deg else "\\sqrt{") + _oseq(_ochild(el, "e")) + "}"
    if tag == "bar":
        return "\\overline{" + _oseq(_ochild(el, "e")) + "}"
    if tag == "acc":
        return "\\hat{" + _oseq(_ochild(el, "e")) + "}"
    if tag == "func":
        name = _oseq(_ochild(el, "fName")).strip()
        cmd = f"\\{name}" if name in ("min", "max", "ln", "log", "exp", "sin", "cos", "tan", "lim") \
            else f"\\operatorname{{{name}}}"
        return cmd + _oseq(_ochild(el, "e"))
    if tag in ("box", "borderBox", "groupChr", "limLow", "limUpp", "m"):
        return "".join(omml_to_latex(c) for c in el if not c.tag.endswith("Pr"))
    if tag.endswith("Pr") or tag in ("ctrlPr",):
        return ""
    return "".join(omml_to_latex(c) for c in el)


def _tidy_latex(latex: str) -> str:
    # Slide callout numbers inside math: "(1)BE_{MAT,y} + (2)BE_{ENE,y}".
    latex = re.sub(r"(?:\\left)?\((?:[1-9]|[A-C])(?:\\right)?\)\s*(?=[A-Za-z{\\])", "", latex)
    latex = re.sub(r"\{(\\overline\{[A-Za-z]+\})\}_", r"\1_", latex)   # {\overline{DP}}_{BL,a}
    # {BE}_{y} -> BE_{y}; Word may also split DHC_{PJ,j} into "DH" + {C}_{PJ,j}, which rejoins here.
    latex = re.sub(r"(?<!\\overline)\{([A-Za-z]+)\}_", r"\1_", latex)
    latex = re.sub(r"\{(\\[a-zA-Z]+ ?)\}_", r"\1_", latex)
    latex = re.sub(r"\{([A-Za-z0-9]+)\}\^", r"\1^", latex)
    latex = re.sub(r"\s+", " ", latex).strip()
    latex = re.sub(r"_\{\s*([^{}]*?)\s*\}", lambda m: "_{" + m.group(1).replace(" ", "") + "}", latex)
    return latex


def _inline(el: ET.Element, ns_t: str) -> str:
    """Text of a paragraph with each math object as $shorthand$ in place."""
    from math_shorthand import _normalize
    out: list[str] = []

    def walk(node: ET.Element) -> None:
        tag = node.tag
        if tag in (M + "oMath",):
            latex = _tidy_latex(omml_to_latex(node))
            out.append(f"${to_shorthand(latex)}$" if latex else "")
            return
        if tag == ns_t:
            out.append(_normalize(node.text or ""))
        for c in node:
            walk(c)
    walk(el)
    return re.sub(r"[ \t]+", " ", "".join(out)).strip()


HEADING = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(.{1,80})$")
# "2.3.4. 온실가스 감축량(ER_y)본 방법론에서 …": heading and its first sentence in one paragraph.
HEADING_MATH = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+([^\x00()]{1,40}?)\s*\(\s*\x00\s*\)")
# "1)  고정 데이터 및 인자", "가. 사업 배출량" (business-plan sub-headings)
HEADING_SUB = re.compile(r"^(?:\d{1,2}\)|[가-하]\.)\s+(.{2,40})$")
EQ_TEXT_OK = re.compile(r"^[\s+\-−×=(),.:]*$")
SLIDE_LABEL = re.compile(r"배출량|감축량|평가\s*범위|데이터\s*및\s*(?:인자|변수)|모니터링")
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _slide_label(text: str) -> str:
    """A slide box that names what follows ("[ (A) 베이스라인 배출량 () ]" -> "베이스라인 배출량")."""
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"^[\[【]\s*|\s*[\]】]$", "", text)
    text = re.sub(r"^\(?[A-Ca-c1-9]\)\s*", "", text)
    text = re.sub(r"\(\s*\)", "", text).strip()
    if not text or len(text) > 32 or re.search(r"[=×*]", text) or text.startswith("본 "):
        return ""
    return text if SLIDE_LABEL.search(text) else ""


def _para_parts(p: ET.Element, ns_t: str) -> tuple[str, list[str]]:
    texts, maths = [], []
    def walk(node: ET.Element) -> None:
        if node.tag == M + "oMath":
            maths.append(_tidy_latex(omml_to_latex(node)))
            texts.append("\x00")
            return
        if node.tag == ns_t:
            texts.append(node.text or "")
        for c in node:
            walk(c)
    walk(p)
    return "".join(texts), maths


def _row_cells(tr: ET.Element, ns_t: str, para_tag: str, cell_tag: str) -> list[dict[str, Any]]:
    cells = []
    for tc in tr.findall(cell_tag):
        paras = [_inline(p, ns_t) for p in tc.iter(para_tag)]
        maths = [_tidy_latex(omml_to_latex(m)) for m in tc.iter(M + "oMath")]
        text_only = " ".join(t for t in ("".join(x.text or "" for x in tc.iter(ns_t)),) if t).strip()
        cells.append({"text": "\n".join(p for p in paras if p), "plain": re.sub(r"\s+", " ", text_only),
                      "maths": maths})
    return cells


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _absorb_table(rows: list[list[dict[str, Any]]], heading: str, found: dict[str, dict[str, Any]],
                  order: list[str]) -> None:
    def upsert(latex: str, **fields: str) -> None:
        try:
            key = param_key(latex)
        except EquationError:
            return
        entry = found.setdefault(key, {"symbol": to_shorthand(key)})
        if key not in order:
            order.append(key)
        for field, value in fields.items():
            value = (value or "").strip()
            if value and value != "-" and not entry.get(field):
                entry[field] = value
    if not rows:
        return
    header = [_norm(c["plain"]) for c in rows[0]]
    if rows[0] and _norm(rows[0][0]["plain"]) in ("데이터/인자", "데이터/변수"):
        labels = {_norm(label): field for form in FORMS.values() for fields in form["fields"].values()
                  for field, label in fields}
        current: Optional[str] = None
        kind = "fixed" if "고정" in heading or "사전" in heading else "monitoring" if "모니터링" in heading else ""
        for row in rows:
            if len(row) < 2:
                continue
            label = _norm(row[0]["plain"])
            if label in ("데이터/인자", "데이터/변수"):
                current = row[1]["maths"][0] if row[1]["maths"] else None
                if current:
                    upsert(current, kind=kind)
                continue
            if current and label in labels:
                upsert(current, **{labels[label]: row[1]["text"]})
        return
    if "기호" in header:
        col = {name: header.index(name) for name in header}
        sym_i = col["기호"]
        name_i = next((col[k] for k in ("정의", "항목", "설명") if k in col), None)
        unit_i = col.get("단위")
        value_i = col.get("값")
        kind_i = next((i for i, h in enumerate(header) if "고정" in h and "모니터링" in h), None)
        remark_i = col.get("비고")
        for row in rows[1:]:
            if sym_i >= len(row) or not row[sym_i]["maths"]:
                continue
            get = lambda i: row[i]["text"] if i is not None and i < len(row) else ""  # noqa: E731
            kind = KIND_ALIASES.get(_norm(get(kind_i)), "") if kind_i is not None else ""
            upsert(row[sym_i]["maths"][0], name=get(name_i), unit=get(unit_i), value=get(value_i),
                   kind=kind, remark=get(remark_i))


def import_office(path: Path) -> dict[str, Any]:
    """Read equations, symbol tables and data/parameter tables from a .docx or .pptx."""
    plan = new_plan()
    plan["title"] = path.stem
    found: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    sections: list[dict[str, Any]] = []
    seen_eq: set[str] = set()
    heading = ""
    current: Optional[dict[str, Any]] = None
    last_eq: Optional[dict[str, Any]] = None
    depth: dict[str, int] = {}

    def section_for(name: str) -> dict[str, Any]:
        nonlocal current
        if current is None or current["name"] != (name or "산정식"):
            current = {"id": f"s{len(sections) + 1}", "name": name or "산정식", "depth": depth.get(name, 1),
                       "description": "", "equations": []}
            sections.append(current)
        return current

    def handle_para(p: ET.Element, ns_t: str) -> None:
        nonlocal heading, last_eq
        text, maths = _para_parts(p, ns_t)
        plain = text.replace("\x00", "").strip()
        head = HEADING_MATH.match(text.strip()) if maths else None
        m = HEADING.match(plain) if not maths and len(plain) < 90 else None
        if head or m:
            number, name = (head or m).group(1), (head or m).group(2)
            heading = re.sub(r"\(\s*\)", "", name.replace("\x00", "")).strip()
            depth[heading] = number.count(".") + 1
            if not maths:
                return
        sub_head = HEADING_SUB.match(plain) if not maths else None
        if sub_head:
            heading = sub_head.group(1).strip()
            depth.setdefault(heading, 9)
            return
        if EQ_TEXT_OK.match(plain):
            pieces = text.split("\x00")
            latex = "".join(_mtext(pieces[i]) + (maths[i] if i < len(maths) else "")
                            for i in range(len(pieces)))
            latex = _tidy_latex(latex)
            if not latex or "=" not in latex and not (last_eq and re.match(r"^[+\-]", latex)):
                return
            if "=" not in latex and last_eq is not None:
                last_eq["text"] = to_shorthand(_tidy_latex(to_latex(last_eq["text"]) + " " + latex))
                return
            if latex in seen_eq:
                return
            seen_eq.add(latex)
            sec = section_for(heading)
            eq = {"id": f"e{len(seen_eq)}", "text": to_shorthand(latex), "label": ""}
            sec["equations"].append(eq)
            last_eq = eq
            return
        # "· ⟦n0⟧ : 통계적으로 필요한 샘플 수" definition bullets
        m = re.match(r"^[·•\-\s]*\x00\s*[:：]\s*(.+)$", text.strip())
        if m and len(maths) == 1:
            try:
                key = param_key(maths[0])
            except EquationError:
                return
            entry = found.setdefault(key, {"symbol": to_shorthand(key)})
            entry.setdefault("name", m.group(1).replace("\x00", "").strip())
            if key not in order:
                order.append(key)

    with zipfile.ZipFile(path) as z:
        if path.suffix.lower() == ".docx":
            root = ET.fromstring(z.read("word/document.xml"))
            body = root.find(W + "body")
            for block in list(body) if body is not None else []:
                if block.tag == W + "p":
                    handle_para(block, W + "t")
                    plain = "".join(t.text or "" for t in block.iter(W + "t")).strip()
                    if HEADING.match(plain) and not list(block.iter(M + "oMath")):
                        last_eq = None
                elif block.tag == W + "tbl":
                    rows = [_row_cells(tr, W + "t", W + "p", W + "tc") for tr in block.findall(W + "tr")]
                    _absorb_table(rows, heading, found, order)
                    last_eq = None
        else:
            names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                           key=lambda n: int(re.findall(r"\d+", n)[-1]))
            for name in names:
                root = ET.fromstring(z.read(name))
                # Shapes and tables top to bottom; short label boxes ("평가 범위 1 : …",
                # "베이스라인 고정데이터 및 인자") name the equations and tables below them.
                items = []
                for el in root.iter():
                    if el.tag in (P + "sp", P + "graphicFrame"):
                        off = el.find(f".//{A}off")
                        y = int(off.get("y", 0)) if off is not None else 0
                        x = int(off.get("x", 0)) if off is not None else 0
                        items.append((y, x, el))
                slide_head = ""
                for _, _, el in sorted(items, key=lambda t: (t[0], t[1])):
                    if el.tag == P + "graphicFrame":
                        for tbl in el.iter(A + "tbl"):
                            rows = [_row_cells(tr, A + "t", A + "p", A + "tc") for tr in tbl.findall(A + "tr")]
                            _absorb_table(rows, f"{slide_head} {heading}", found, order)
                        continue
                    label = _slide_label("".join(t.text or "" for t in el.iter(A + "t")))
                    if label and not list(el.iter(M + "oMath")):
                        heading = label
                        slide_head = slide_head or label
                        depth.setdefault(label, 1)
                        last_eq = None
                        continue
                    for p in el.iter(A + "p"):
                        handle_para(p, A + "t")
    shallow = min((s["depth"] for s in sections), default=1)
    for sec in sections:
        sec["level"] = 1 if sec.pop("depth") <= shallow else 2
    plan["sections"] = sections
    plan["parameters"] = [found[k] for k in order]
    return plan


# ---------------------------------------------------------------------------
# Word output (KCS layout)
# ---------------------------------------------------------------------------

def _docx_rich(text: str, font: str, size: float, bold: bool = False, color: Optional[str] = None) -> str:
    from office_equations import _w_run
    from omml_math import docx_rpr, omml_math
    out, pos = [], 0
    for m in INLINE_MATH.finditer(text or ""):
        if m.start() > pos:
            out.append(_w_run(text[pos:m.start()], font, size, bold, color))
        try:
            out.append(omml_math(to_latex(m.group(1)), docx_rpr(size, color=color, ea_font=font)))
        except EquationError:
            out.append(_w_run(m.group(0), font, size, bold, color))
        pos = m.end()
    if pos < len(text or ""):
        out.append(_w_run(text[pos:], font, size, bold, color))
    return "".join(out) or _w_run("", font, size)


def _docx_cell_paras(text: str, font: str, size: float, align: str = "left", bold: bool = False,
                     color: Optional[str] = None) -> str:
    lines = (text or "").split("\n") or [""]
    return "".join(f'<w:p><w:pPr><w:spacing w:before="0" w:after="0"/><w:jc w:val="{align}"/></w:pPr>'
                   f"{_docx_rich(line, font, size, bold, color)}</w:p>" for line in lines)


def build_docx_body(derived: dict[str, Any]) -> list[str]:
    from office_equations import (BAND_FILL, HEADER_FILL, HEADER_TEXT, _table, _tc, _tr, _w_para,
                                  _w_run)
    from omml_math import docx_rpr, omml_math, omml_para
    font, size = derived["font"], derived["font_size_pt"]
    small = size - 1
    rpr_eq = docx_rpr(size + 1.5, ea_font=font)
    rpr_sym = docx_rpr(small, ea_font=font)
    opts = derived["options"]
    form = form_of(opts)
    by_key = {p["key"]: p for p in derived["parameters"]}
    body: list[str] = []
    width = 9500

    def head_cells(labels: list[tuple[str, int]]) -> str:
        return _tr([_tc(_w_run(label, font, small, bold=True, color=HEADER_TEXT), w, fill=HEADER_FILL,
                        align="center") for label, w in labels], header=True)

    if derived.get("title"):
        body.append(_w_para(_w_run(derived["title"], font, size + 5, bold=True, color=HEADER_FILL),
                            align="center", after=240))
    for sec in derived["sections"]:
        level = sec.get("level", 1)
        body.append(_w_para(_w_run(sec.get("name", ""), font, size + (3 if level == 1 else 1.5), bold=True),
                            before=240 if level == 1 else 160, after=100,
                            style="Heading1" if level == 1 else "Heading2", keep_next=True))
        if sec.get("description"):
            body.append(_w_para(_docx_rich(sec["description"], font, size), after=100, keep_next=True))
        for eq in sec["equations"]:
            if not eq.get("latex"):
                continue
            eq_par = (f'<w:p><w:pPr><w:spacing w:before="60" w:after="60"/><w:jc w:val="center"/></w:pPr>'
                      f"{omml_para(eq['latex'], rpr_eq)}</w:p>")
            label = _w_run(f"({eq['label']})" if eq.get("label") else "", font, size)
            body.append(_table([_tr([_tc(eq_par, 8300, borders=False),
                                     _tc(label, 1200, align="right", borders=False)])],
                               [8300, 1200], borders=False))
        if sec["symbols"] and form["symbol_style"] == "where":
            # "여기서," list: symbol : definition [unit], borderless.
            widths = [1900, 300, 5500, 1800]
            rows = []
            for key in sec["symbols"]:
                p = by_key[key]
                unit = f"[{p['unit']}]" if p.get("unit") else ""
                rows.append(_tr([_tc(omml_math(p["latex"], rpr_sym), widths[0], align="right", borders=False),
                                 _tc(_w_run(":", font, size), widths[1], align="center", borders=False),
                                 _tc(_docx_cell_paras(p.get("name", ""), font, size), widths[2], borders=False),
                                 _tc(_docx_cell_paras(unit, font, size), widths[3], borders=False)]))
            body.append(_w_para(_w_run("여기서,", font, size), before=60, after=40, keep_next=True))
            body.append(_table(rows, widths, borders=False))
        elif sec["symbols"]:
            show_value = opts.get("show_values") and any(by_key[k].get("value") for k in sec["symbols"])
            cols = [("기호", 1900), ("정의", 5000 if show_value else 6000)] + \
                   ([("값", 1100)] if show_value else []) + [("단위", 1500 if show_value else 1600)]
            rows = [head_cells(cols)]
            for i, key in enumerate(sec["symbols"]):
                p = by_key[key]
                fill = BAND_FILL if i % 2 else None
                cells = [_tc(omml_math(p["latex"], rpr_sym), cols[0][1], fill=fill, align="center"),
                         _tc(_docx_cell_paras(p.get("name", ""), font, small), cols[1][1], fill=fill)]
                if show_value:
                    cells.append(_tc(_docx_cell_paras(p.get("value", ""), font, small, "center"),
                                     cols[2][1], fill=fill))
                cells.append(_tc(_docx_cell_paras(p.get("unit", ""), font, small, "center"),
                                 cols[-1][1], fill=fill))
                rows.append(_tr(cells))
            body.append(_w_para("", after=40))
            body.append(_table(rows, [w for _, w in cols]))
            body.append(_w_para("", after=80))

    params = [p for p in derived["parameters"] if p["kind"] in ("monitoring", "fixed")]
    if params:
        body.append(_w_para(_w_run(form["chapter"], font, size + 3, bold=True), before=280, after=100,
                            style="Heading1", keep_next=True))
        body.append(_w_para(_w_run(form["procedure"], font, size + 1.5, bold=True), before=120, after=80,
                            style="Heading2", keep_next=True))
        body.append(_w_para(_w_run(form["procedure_lead"], font, size), after=80, keep_next=True))
        cols = [(label, w) for _, label, w in form["summary"]]
        rows = [head_cells(cols)]
        for i, p in enumerate(params, start=1):
            fill = BAND_FILL if i % 2 == 0 else None
            cells = []
            for (field, _, w) in form["summary"]:
                if field == "symbol":
                    cells.append(_tc(omml_math(p["latex"], rpr_sym), w, fill=fill, align="center"))
                    continue
                value = str(i) if field == "no" else form["kind_names"][p["kind"]] if field == "kind" \
                    else p.get(field, "")
                cells.append(_tc(_docx_cell_paras(value, font, small, "left" if field == "name" else "center"),
                                 w, fill=fill))
            rows.append(_tr(cells))
        body.append(_table(rows, [w for _, w in cols]))
        for kind, title, lead in form["groups"]:
            group = [p for p in params if p["kind"] == kind]
            if not group:
                continue
            body.append(_w_para(_w_run(title, font, size + 1.5, bold=True), before=200, after=80,
                                style="Heading2", keep_next=True))
            body.append(_w_para(_w_run(lead, font, size), after=80, keep_next=True))
            for p in group:
                rows = [_tr([_tc(_w_run(form["param_word"], font, small, bold=True), 2200, fill=BAND_FILL,
                                 align="center"),
                             _tc(omml_math(p["latex"], rpr_sym), width - 2200, align="center")])]
                for field, label in form["fields"][kind]:
                    value = p.get(field) or (p.get("purpose_auto", "") if field == "purpose" else "")
                    rows.append(_tr([_tc(_w_run(label, font, small, bold=True), 2200, fill=BAND_FILL,
                                         align="center"),
                                     _tc(_docx_cell_paras(value or "-", font, small), width - 2200)]))
                # Keep one parameter's table on one page: every row but the last keeps with the next.
                rows[:-1] = [r.replace("<w:pPr><w:spacing", "<w:pPr><w:keepNext/><w:spacing") for r in rows[:-1]]
                body.append(_table(rows, [2200, width - 2200]))
                body.append(_w_para("", after=120))
    return body


def write_docx(derived: dict[str, Any], out: Path) -> None:
    from office_equations import W_NS, _settings_xml, _styles_xml, require_new_output
    from omml_math import M_NS
    require_new_output(out)
    body = "".join(build_docx_body(derived))
    sect = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1200" '
            'w:bottom="1440" w:left="1200" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>')
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f'<w:document xmlns:w="{W_NS}" xmlns:m="{M_NS}" '
                f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                f"<w:body>{body}{sect}</w:body></w:document>")
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
            '<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
            '</Types>'),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>'),
        "word/_rels/document.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/>'
            '</Relationships>'),
        "word/document.xml": document,
        "word/styles.xml": _styles_xml(derived["font"], derived["font_size_pt"]),
        "word/settings.xml": _settings_xml(),
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "x", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


# ---------------------------------------------------------------------------
# PowerPoint output
# ---------------------------------------------------------------------------

def write_pptx(derived: dict[str, Any], out: Path) -> int:
    """Build slides with python-pptx, then turn {{eq:}}/{{sym:}} tokens into OMML."""
    import tempfile

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    from office_equations import inject_pptx, require_new_output
    from omml_math import equation_height_factor, estimate_width_em
    require_new_output(out)
    font = derived["font"]
    navy, band, white, ink = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0xF3, 0xF5, 0xF8), \
        RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x22, 0x22, 0x22)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    inline: dict[str, str] = {}
    left, usable_w, top0, bottom = Inches(0.6), Inches(12.13), Inches(1.15), Inches(7.0)

    def tokens(text: str) -> str:
        def repl(m: re.Match) -> str:
            try:
                latex = to_latex(m.group(1))
            except EquationError:
                return m.group(0)
            tid = f"m{len(inline) + 1}"
            inline[tid] = latex
            return f"{{{{eq:{tid}}}}}"
        return INLINE_MATH.sub(repl, text or "")

    def style_run(run, size: float, bold: bool = False, color=ink) -> None:
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
        run.font.color.rgb = color
        rpr = run._r.get_or_add_rPr()
        ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": font})
        rpr.append(ea)

    def new_slide(title: str):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(left, Inches(0.35), usable_w, Inches(0.6))
        tf = box.text_frame
        tf.word_wrap = True
        style_run(tf.paragraphs[0].add_run(), 22, True, navy)
        tf.paragraphs[0].runs[0].text = title
        line = slide.shapes.add_connector(1, left, Inches(0.98), left + usable_w, Inches(0.98))
        line.line.color.rgb = navy
        return slide

    def set_cell(cell, text: str, size: float, bold: bool = False, color=ink, fill=None,
                 align=PP_ALIGN.LEFT) -> None:
        cell.text = ""
        tf = cell.text_frame
        tf.word_wrap = True
        for i, line in enumerate((text or "").split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = line
            style_run(run, size, bold, color)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        if fill is not None:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
        else:
            cell.fill.background()

    def add_table(slide, top: int, header: list[tuple[str, float]], rows: list[list[str]],
                  aligns: list, row_h: float = 0.36, size: float = 12, width=None, x=None):
        table_w = width or usable_w
        shape = slide.shapes.add_table(len(rows) + (1 if header[0][0] else 0), len(header),
                                       x if x is not None else left, top, table_w,
                                       Inches(row_h * (len(rows) + 1)))
        tbl = shape.table
        total = sum(w for _, w in header)
        for i, (_, w) in enumerate(header):
            tbl.columns[i].width = Emu(int(table_w * w / total))
        r0 = 0
        if header[0][0]:
            for i, (label, _) in enumerate(header):
                set_cell(tbl.cell(0, i), label, size, True, white, navy, PP_ALIGN.CENTER)
            r0 = 1
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                fill = band if (r % 2 and header[0][0]) or (not header[0][0] and c == 0) else None
                bold = not header[0][0] and c == 0
                set_cell(tbl.cell(r + r0, c), value, size, bold, ink, fill, aligns[c])
        for r in range(len(rows) + r0):
            tbl.rows[r].height = Inches(row_h)
        tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
        tbl_pr.set("firstRow", "0")
        tbl_pr.set("bandRow", "0")
        return shape

    by_key = {p["key"]: p for p in derived["parameters"]}
    C, L = PP_ALIGN.CENTER, PP_ALIGN.LEFT
    for sec in derived["sections"]:
        eqs = [e for e in sec["equations"] if e.get("latex")]
        sym_rows = [[f"{{{{sym:{by_key[k]['id']}}}}}", tokens(by_key[k].get("name", "")),
                     by_key[k].get("unit", "")] for k in sec["symbols"]]
        if not eqs and not sym_rows:
            continue
        slide = new_slide(sec.get("name", ""))
        y = top0
        if sec.get("description"):
            box = slide.shapes.add_textbox(left, y, usable_w, Inches(0.5))
            box.text_frame.word_wrap = True
            run = box.text_frame.paragraphs[0].add_run()
            run.text = tokens(sec["description"])
            style_run(run, 13)
            y += Inches(0.55)
        for eq in eqs:
            # Shrink long equations to the box width (11.1 in ≈ 800 pt), never below 10 pt.
            size = max(10.0, min(18.0, 0.9 * (usable_w - Inches(1.0)) / Pt(1) / estimate_width_em(eq["latex"])))
            h = Inches(0.3 + 0.28 * equation_height_factor(eq["latex"]) * size / 18)
            box = slide.shapes.add_textbox(left, y, usable_w - Inches(1.0), h)
            para = box.text_frame.paragraphs[0]
            para.alignment = C
            run = para.add_run()
            run.text = f"{{{{eq:{eq['id']}}}}}"
            style_run(run, round(size, 1))
            if eq.get("label"):
                lab = slide.shapes.add_textbox(left + usable_w - Inches(1.0), y, Inches(1.0), h)
                lab.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                lab.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                r = lab.text_frame.paragraphs[0].add_run()
                r.text = f"({eq['label']})"
                style_run(r, 13)
            y += h + Inches(0.08)
        header = [("기호", 2.2), ("정의", 7.4), ("단위", 2.4)]
        while sym_rows:
            room = max(1, int((bottom - y) / Inches(0.4)) - 1)
            chunk, sym_rows = sym_rows[:room], sym_rows[room:]
            add_table(slide, y + Inches(0.1), header, chunk, [C, L, C], 0.4)
            if sym_rows:
                slide = new_slide(f"{sec.get('name', '')} (계속)")
                y = top0

    params = [p for p in derived["parameters"] if p["kind"] in ("monitoring", "fixed")]
    if params:
        form = form_of(derived["options"])

        def summary_cell(field: str, i: int, p: dict[str, Any]) -> str:
            if field == "no":
                return str(i)
            if field == "symbol":
                return f"{{{{sym:{p['id']}}}}}"
            if field == "kind":
                return form["kind_names"][p["kind"]]
            return tokens(p.get(field, ""))
        rows = [[summary_cell(f, i, p) for f, _, _ in form["summary"]] for i, p in enumerate(params, start=1)]
        header = [(label, w / 1000) for _, label, w in form["summary"]]
        aligns = [L if f == "name" else C for f, _, _ in form["summary"]]
        page = 13
        for start in range(0, len(rows), page):
            slide = new_slide(form["procedure"] + (" (계속)" if start else ""))
            add_table(slide, top0, header, rows[start:start + page], aligns, 0.42)
        for kind, title, _ in form["groups"]:
            group = [p for p in params if p["kind"] == kind]
            half = (usable_w - Inches(0.3)) // 2
            for start in range(0, len(group), 2):
                slide = new_slide(title)
                for j, p in enumerate(group[start:start + 2]):
                    detail = [[form["param_word"], f"{{{{sym:{p['id']}}}}}"]]
                    for field, label in form["fields"][kind]:
                        value = p.get(field) or (p.get("purpose_auto", "") if field == "purpose" else "")
                        detail.append([label, tokens(value) or "-"])
                    add_table(slide, top0, [("", 1.0), ("", 2.6)], detail, [C, L], 0.5, 10.5,
                              width=half, x=left + j * (half + Inches(0.3)))

    with tempfile.TemporaryDirectory() as tmp:
        draft = Path(tmp) / "draft.pptx"
        prs.save(draft)
        spec = to_equation_spec(derived)
        for tid, latex in inline.items():
            spec["scenarios"].append({"id": f"_{tid}", "equations": [{"id": tid, "latex": latex}]})
        spec["monitoring_parameters"] = spec["monitoring_parameters"] + spec.pop("calculated_parameters")
        count, problems = inject_pptx(draft, spec, out)
        if problems:
            raise EquationError("; ".join(problems))
    return count


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

EXAMPLE_TEXT = """# 베이스라인 배출량
> 사업이 없을 때 기존 설비가 계통전력으로 같은 양을 공급한 경우의 배출량
BEy = EG_PJ,y × EF_grid,y

# 사업 배출량
PEy = sum_i FC_i,y × NCV_i × EF_CO2,i

# 감축량
ERy = BEy - PEy - LEy

EG_PJ,y | y년도 사업 순발전량 | MWh/년 | 모니터링
EF_grid,y | y년도 계통 배출계수 | tCO2-eq/MWh | 고정 | 0.4173
FC_i,y | y년도 연료 i 사용량 | kL/년 | 모니터링
NCV_i | 연료 i의 순발열량 | TJ/kL | 고정
EF_CO2,i | 연료 i의 CO2 배출계수 | tCO2/TJ | 고정
BEy | y년도 베이스라인 배출량 | tCO2-eq/년
PEy | y년도 사업 배출량 | tCO2-eq/년
LEy | y년도 누출량 | tCO2-eq/년
ERy | y년도 온실가스 감축량 | tCO2-eq/년
"""


def load_plan(path: Path) -> dict[str, Any]:
    if path.suffix.lower() in (".txt", ".md"):
        return plan_from_text(path.read_text(encoding="utf-8"))
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(data: Any, out: Optional[Path]) -> None:
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text + "\n", encoding="utf-8")
        print(f"OK  {out}")
    else:
        print(text)


def _print_issues(derived: dict[str, Any]) -> int:
    errors = 0
    for issue in derived["issues"]:
        if issue["level"] == "info":
            continue
        errors += issue["level"] == "error"
        print(f"{issue['level'].upper():5}  {issue['where']}: {issue['msg']}")
    eqs = sum(1 for s in derived["sections"] for e in s["equations"] if e.get("latex"))
    kinds = {k: sum(1 for p in derived["parameters"] if p["kind"] == k) for k in KINDS}
    print(f"equations={eqs} parameters={len(derived['parameters'])} "
          f"(산출 {kinds['calculated']}, 모니터링 {kinds['monitoring']}, 고정 {kinds['fixed']}) errors={errors}")
    return errors


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Reduction-project monitoring plan: rough-text equations "
                                     "-> editable OMML equations and KCS tables.")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("example", help="print a sample plan")
    ft = sub.add_parser("from-text", help="quick text -> plan JSON")
    ft.add_argument("text", type=Path)
    ft.add_argument("-o", "--output", type=Path)
    im = sub.add_parser("import", help="read equations/tables from .docx or .pptx")
    im.add_argument("file", type=Path)
    im.add_argument("-o", "--output", type=Path)
    ck = sub.add_parser("check", help="derive and list issues")
    ck.add_argument("plan", type=Path)
    for name in ("docx", "pptx", "spec"):
        p = sub.add_parser(name)
        p.add_argument("plan", type=Path)
        p.add_argument("-o", "--output", type=Path, required=name != "spec")
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    from console_encoding import configure_utf8_stdio
    configure_utf8_stdio()
    args = build_parser().parse_args(argv)
    try:
        if args.command == "example":
            _write_json(plan_from_text(EXAMPLE_TEXT), None)
            return 0
        if args.command == "from-text":
            _write_json(plan_from_text(args.text.read_text(encoding="utf-8")), args.output)
            return 0
        if args.command == "import":
            plan = import_office(args.file)
            _write_json(plan, args.output)
            _print_issues(derive(plan))
            return 0
        derived = derive(load_plan(args.plan))
        errors = _print_issues(derived)
        if args.command == "check":
            return 1 if errors else 0
        if errors:
            print("FAIL  fix the errors above first", file=sys.stderr)
            return 1
        if args.command == "spec":
            _write_json(to_equation_spec(derived), args.output)
        elif args.command == "docx":
            write_docx(derived, args.output)
            print(f"OK  {args.output}")
        else:
            count = write_pptx(derived, args.output)
            print(f"OK  {args.output}  equations={count}")
    except (EquationError, FileExistsError) as exc:
        print(f"ERROR  {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
